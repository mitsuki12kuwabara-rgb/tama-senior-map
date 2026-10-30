# -*- coding: utf-8 -*-
"""
多摩市 高齢者施設マップ  data.js 再生成スクリプト

使い方（施設の追加・修正をしたあと）:
    1. 元Excel（tama_senior_facilities_YYYYMMDD.xlsx）の「施設一覧」シートを編集
    2. このフォルダで:  python -I regenerate.py  [Excelのパス]
       （パス省略時は ../data/ 内の .xlsx、無ければデスクトップの原本を探します）
    3. ひとつ上の data.js が更新される → git add/commit/push

住所から緯度経度は国土地理院ジオコーダ(無料・商用可)で取得し tools/geocache.json に保存。
二回目以降はキャッシュを使うため再取得しません（新規住所のみ取得）。
インタビュー状況は地図画面(端末保存)で記録し、CSVエクスポート→Excelへ転記します。
"""
import openpyxl, json, os, sys, math, time, urllib.request, urllib.parse
from collections import defaultdict, Counter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
CACHE = os.path.join(HERE, "geocache.json")

def find_xlsx():
    if len(sys.argv) > 1 and os.path.exists(sys.argv[1]):
        return sys.argv[1]
    data_dir = os.path.join(ROOT, "data")
    if os.path.isdir(data_dir):
        for f in sorted(os.listdir(data_dir)):
            if f.lower().endswith(".xlsx") and not f.startswith("~"):
                return os.path.join(data_dir, f)
    # fallback: Desktop original
    desk = os.path.join(os.path.expanduser("~"), "OneDrive", "デスクトップ", "高齢者絵本")
    if os.path.isdir(desk):
        for f in sorted(os.listdir(desk)):
            if f.lower().endswith(".xlsx") and "backup" not in f and not f.startswith("~"):
                return os.path.join(desk, f)
    raise SystemExit("Excelが見つかりません。パスを引数で指定してください。")

SRC = find_xlsx()
print("source:", SRC)
wb = openpyxl.load_workbook(SRC, data_only=True)
names = wb.sheetnames
ws_overview, ws_list, ws_consult = wb[names[0]], wb[names[1]], wb[names[2]]

COLS = ["no","category","name","address","tel","actual","actual_type","actual_asof",
        "capacity","unit","capacity_type","area","access","notes","src_name","src_num",
        "carepro_update","survey_date"]

facilities = []
seen_header = False
for row in ws_list.iter_rows(values_only=True):
    if row[0] == "No.": seen_header = True; continue
    if not seen_header or row[0] is None: continue
    facilities.append({k: (row[i] if i < len(row) else None) for i, k in enumerate(COLS)})
print("facilities:", len(facilities))

# ---- geocode ----
cache = json.load(open(CACHE, encoding="utf-8")) if os.path.exists(CACHE) else {}
def clean(a): return a.replace("　"," ").strip().split(" ")[0] if a else a
def geocode(addr):
    if not addr: return None
    if addr in cache: return cache[addr]
    for attempt in (clean(addr), addr):
        try:
            url = "https://msearch.gsi.go.jp/address-search/AddressSearch?q=" + urllib.parse.quote(attempt)
            d = json.load(urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent":"Mozilla/5.0"}), timeout=25))
            if d:
                c = d[0]["geometry"]["coordinates"]
                cache[addr] = {"lat": c[1], "lng": c[0]}; time.sleep(0.25); return cache[addr]
        except Exception as e:
            sys.stderr.write("geo err %s: %s\n" % (attempt, e))
        time.sleep(0.3)
    cache[addr] = None; return None

for f in facilities:
    g = geocode(f["address"])
    f["lat"] = g["lat"] if g else None
    f["lng"] = g["lng"] if g else None
    if not g: print("  GEOCODE FAIL:", f["no"], f["name"], f["address"])

# consult windows
consult = []
start = False
for row in ws_consult.iter_rows(values_only=True):
    r = list(row)
    if r and r[0] == "窓口名": start = True; continue
    if start and r and r[0]:
        g = geocode(r[1])
        consult.append({"name": r[0], "address": r[1], "tel": r[2], "note": r[3], "src": r[4],
                        "lat": g["lat"] if g else None, "lng": g["lng"] if g else None})
json.dump(cache, open(CACHE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# ---- groups ----
def group_of(cat):
    if "老人福祉館" in cat or "介護予防" in cat: return "koryu"
    if any(k in cat for k in ("通所","デイ","多機能","リハビリ")): return "tsusho"
    return "sumai"
GROUPS = [
    {"id":"sumai","label":"住まい（入居系）","color":"#e8622c"},
    {"id":"tsusho","label":"通所・多機能","color":"#2f74c0"},
    {"id":"koryu","label":"交流・介護予防","color":"#2e9e5b"},
    {"id":"madoguchi","label":"相談・連携の窓口","color":"#8457c4"},
]
for f in facilities: f["group"] = group_of(f["category"])
print("group counts:", dict(Counter(f["group"] for f in facilities)))

# ---- jitter duplicate coords ----
by_coord = defaultdict(list)
for i, f in enumerate(facilities):
    if f["lat"] is not None: by_coord[(round(f["lat"],6), round(f["lng"],6))].append(i)
R = 0.00022
for coord, idxs in by_coord.items():
    if len(idxs) == 1:
        f = facilities[idxs[0]]; f["dlat"], f["dlng"] = f["lat"], f["lng"]; continue
    n = len(idxs)
    for k, i in enumerate(idxs):
        ang = 2*math.pi*k/n; r = R*(1 + (n>6)*0.6)
        facilities[i]["dlat"] = facilities[i]["lat"] + r*math.cos(ang)
        facilities[i]["dlng"] = facilities[i]["lng"] + r*math.sin(ang)/math.cos(math.radians(facilities[i]["lat"]))

# ---- meta ----
ov = [list(r) for r in ws_overview.iter_rows(values_only=True)]
caution_labels = ["実人数の読み方","時点","定員・戸数","重複","未確認","網羅性","対象外","資料の優先","閉館予定","駅エリア","連絡時の確認"]
cautions = [{"label": r[0], "text": r[1]} for r in ov if r and r[0] in caution_labels and r[1]]
# 駅近で検討する候補（概要シート col3-6）と補足メモ
station = [{"name":r[3],"walk":r[4],"tel":r[5],"people":r[6]}
           for r in ov if r and len(r) > 6 and r[3] and r[4] and r[5] and r[3] != "駅周辺で検討する例"]
station_names = [s["name"] for s in station]
station_notes = [r[3] for r in ov if r and len(r) > 3 and r[3]
                 and r[3] not in station_names and r[3] != "駅周辺で検討する例"]
gc = Counter(f["group"] for f in facilities)
meta = {
  "title":"多摩市 高齢者施設マップ",
  "subtitle":"高齢者の人生を絵本にする活動のための 訪問・インタビュー相談先",
  "surveyDate":"2026年10月29日",
  "purpose":"高齢者の人生を絵本にする活動の、訪問・インタビュー・読み合わせ会の相談先を探す",
  "target":"所在地が東京都多摩市内の入居施設、通所施設、小規模多機能、老人福祉館・介護予防拠点",
  "counts":{"total":len(facilities),"sumai":gc.get("sumai",0),"tsusho":gc.get("tsusho",0),
            "koryu":gc.get("koryu",0),"madoguchi":len(consult),
            "withActual":sum(1 for f in facilities if f["actual"] not in (None,"未確認"))},
  "cautions":cautions,
  "stationExamples":station,
  "stationNotes":station_notes,
}

def norm(v):
    return v.isoformat() if hasattr(v,"isoformat") else v
clean_f = []
for f in facilities:
    clean_f.append({
        "no":f["no"],"group":f["group"],"category":f["category"],"name":f["name"],
        "address":f["address"],"tel":norm(f["tel"]),"actual":norm(f["actual"]),
        "actualType":f["actual_type"],"actualAsof":f["actual_asof"],
        "capacity":norm(f["capacity"]),"unit":f["unit"],"capacityType":f["capacity_type"],
        "area":f["area"],"access":f["access"],"notes":f["notes"],
        "srcName":f["src_name"],"srcNum":f["src_num"],"careproUpdate":norm(f["carepro_update"]),
        "surveyDate":norm(f["survey_date"]),
        "lat":round(f["lat"],6) if f["lat"] else None,"lng":round(f["lng"],6) if f["lng"] else None,
        "dlat":round(f.get("dlat") or f["lat"],6) if f["lat"] else None,
        "dlng":round(f.get("dlng") or f["lng"],6) if f["lng"] else None,
    })
data = {"meta":meta,"groups":GROUPS,"facilities":clean_f,"consult":consult}
out = ("// 多摩市 高齢者施設マップ データ (自動生成: regenerate.py)\n"
       "// 施設の追加・修正はExcelを編集し再生成。インタビュー記録は地図画面(端末保存)→CSVでExcelへ書き戻し\n"
       "window.TAMA_DATA = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n")
open(os.path.join(ROOT, "data.js"), "w", encoding="utf-8").write(out)
print("WROTE", os.path.join(ROOT, "data.js"), "facilities=%d consult=%d" % (len(clean_f), len(consult)))
